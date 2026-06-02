from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from import_goldco_influencer_xlsx_rumble import (  # noqa: E402
    RumbleCandidate,
    _dedupe_existing_channel,
    _best_candidate,
    _choose_rumble_url,
    _extract_rumble_url_from_cell,
    _find_header_maps,
    _iter_source_rows,
    _strict_rumble_candidate,
)


def test_strict_rumble_candidate_accepts_only_channel_paths() -> None:
    assert _strict_rumble_candidate("https://rumble.com/c/DrEricBerg") is not None
    assert _strict_rumble_candidate("https://rumble.com/user/DrEricBerg") is not None
    assert _strict_rumble_candidate("https://rumble.com/playlists") is None
    assert _strict_rumble_candidate("https://rumble.com/DrEricBerg") is None


def test_extract_rumble_url_from_cell_canonicalizes_messy_text() -> None:
    value = "Check https://www.rumble.com/c/DrEricBerg?foo=1, thanks"

    assert _extract_rumble_url_from_cell(value) == "https://rumble.com/c/DrEricBerg"


def test_find_header_maps_and_iter_source_rows() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Q1 HUNT"
    ws["A1"] = "Influencer Name"
    ws["B1"] = "Contact Status"
    ws["C1"] = "Rumble URL"
    ws["A2"] = "Dr. Eric Berg"
    ws["C2"] = "https://rumble.com/c/DrEricBerg"
    ws["A3"] = "Lex Fridman"
    ws["C3"] = ""

    header_map = _find_header_maps(ws)
    assert header_map is not None
    rows = list(_iter_source_rows(ws, header_map))

    assert len(rows) == 2
    assert rows[0].influencer_name == "Dr. Eric Berg"
    assert rows[0].rumble_url_raw == "https://rumble.com/c/DrEricBerg"
    assert rows[1].influencer_name == "Lex Fridman"
    assert rows[1].rumble_url_raw == ""


def test_best_candidate_rejects_ambiguous_close_scores(monkeypatch) -> None:
    monkeypatch.setattr(
        "import_goldco_influencer_xlsx_rumble._resolve_rumble_from_serper_cached",
        lambda name: (
            RumbleCandidate(
                channel_url="https://rumble.com/c/example-a",
                confidence=0.90,
                source="serper",
                query="q1",
            ),
            RumbleCandidate(
                channel_url="https://rumble.com/c/example-b",
                confidence=0.86,
                source="serper",
                query="q2",
            ),
        ),
    )
    assert _best_candidate("Example A") is None


def test_choose_rumble_url_prefers_excel_value(monkeypatch) -> None:
    monkeypatch.setattr(
        "import_goldco_influencer_xlsx_rumble._best_candidate",
        lambda name: RumbleCandidate(
            channel_url="https://rumble.com/c/example",
            confidence=0.95,
            source="serper",
            query="q",
        ),
    )
    row = type("Row", (), {"sheet_name": "Sheet1", "excel_row": 2, "influencer_name": "Example", "rumble_url_raw": "https://rumble.com/c/Direct"})()
    chosen_url, source, confidence, error = _choose_rumble_url(row, 0.88)

    assert chosen_url == "https://rumble.com/c/Direct"
    assert source == "excel"
    assert confidence == 1.0
    assert error is None


def test_dedupe_existing_channel_checks_primary_url_first() -> None:
    class _Response:
        def __init__(self, data):
            self.data = data

    class _Query:
        def __init__(self):
            self._mode = None
            self._value = None

        def select(self, *_args, **_kwargs):
            return self

        def eq(self, column, value):
            self._mode = "eq"
            self._value = (column, value)
            return self

        def contains(self, column, value):
            self._mode = "contains"
            self._value = (column, value)
            return self

        def limit(self, *_args, **_kwargs):
            return self

        def execute(self):
            column, value = self._value
            if self._mode == "eq" and column == "channel_url" and value == "https://rumble.com/c/example":
                return _Response([{"id": "1", "channel_url": value}])
            return _Response([])

    class _Client:
        def table(self, *_args, **_kwargs):
            return _Query()

    assert _dedupe_existing_channel(_Client(), "https://rumble.com/c/example") is not None
    assert _dedupe_existing_channel(_Client(), "https://rumble.com/c/missing") is None


def test_dedupe_existing_channel_checks_secondary_urls() -> None:
    class _Response:
        def __init__(self, data):
            self.data = data

    class _Query:
        def __init__(self):
            self._mode = None
            self._value = None

        def select(self, *_args, **_kwargs):
            return self

        def eq(self, column, value):
            self._mode = "eq"
            self._value = (column, value)
            return self

        def contains(self, column, value):
            self._mode = "contains"
            self._value = (column, value)
            return self

        def limit(self, *_args, **_kwargs):
            return self

        def execute(self):
            column, value = self._value
            if self._mode == "contains" and column == "secondary_urls" and value == ["https://rumble.com/c/secondary"]:
                return _Response([{"id": "2", "channel_url": "https://rumble.com/c/other"}])
            return _Response([])

    class _Client:
        def table(self, *_args, **_kwargs):
            return _Query()

    assert _dedupe_existing_channel(_Client(), "https://rumble.com/c/secondary") is not None
