"""Import Goldco influencer workbook rows into ``channels``.

This is a conservative Rumble-only ingestion utility.

Workflow:
- Read one or more worksheets from an Excel workbook.
- Find the columns named ``Influencer Name`` and ``Rumble URL``.
- If a row already has a Rumble URL, canonicalize and validate it.
- If the URL is missing, search Serper for a precise Rumble channel URL.
- Verify the best Serper candidate before trusting it.
- Skip any URL already present in Supabase ``channels``.
- Default to dry-run and emit a JSON report.

The resolver is intentionally strict:
- Only Rumble channel pages are accepted.
- Only channel-style paths like ``/c/<slug>`` or ``/user/<slug>`` are allowed.
- Ambiguous or weak matches are not inserted.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from functools import lru_cache
from pathlib import Path
from typing import Iterable
from urllib.parse import urlsplit

import httpx
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from postgrest.exceptions import APIError

sys.path.insert(0, str(Path(__file__).resolve().parent))

from core.config import scraper_settings
from core.supabase import get_supabase_client
from utils.channel_urls import canonicalize_channel_url, extract_supported_channel_urls

SERPER_URL = "https://google.serper.dev/search"
SERPER_RETRY_CODES = {408, 425, 429, 500, 502, 503, 504}
SERPER_MAX_ATTEMPTS = 3
SERPER_RETRY_DELAY_SECONDS = 1.5
DEFAULT_SEARCH_LIMIT = 10
DEFAULT_MIN_CONFIDENCE = 0.88


@dataclass(frozen=True)
class SheetHeaderMap:
    sheet_name: str
    header_row: int
    influencer_col: int
    rumble_col: int


@dataclass(frozen=True)
class SourceRow:
    sheet_name: str
    excel_row: int
    influencer_name: str
    rumble_url_raw: str


@dataclass(frozen=True)
class RumbleCandidate:
    channel_url: str
    confidence: float
    source: str
    query: str
    title: str | None = None
    snippet: str | None = None
    display_link: str | None = None
    verified_title: str | None = None
    verified_canonical_url: str | None = None


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _normalize(text: str | None) -> str:
    return " ".join((text or "").strip().split())


def _normalize_for_match(text: str | None) -> str:
    return re.sub(r"[^a-z0-9]+", "", _normalize(text).lower())


def _strip_parenthetical(text: str) -> str:
    return _normalize(re.sub(r"\s*\([^)]*\)\s*", " ", text or ""))


def _split_aliases(raw_name: str) -> list[str]:
    """Generate a short alias list for search queries."""

    base = _normalize(raw_name)
    if not base:
        return []

    aliases: list[str] = []

    def add(value: str) -> None:
        cleaned = _normalize(value)
        if cleaned and cleaned not in aliases:
            aliases.append(cleaned)

    add(base)
    add(_strip_parenthetical(base))
    add(base.replace("_", " "))

    if ";" in base or "|" in base:
        for part in re.split(r"[;|]", base):
            add(part)

    if " / " in base:
        for part in base.split(" / "):
            add(part)

    if " - " in base:
        parts = [part.strip() for part in base.split(" - ") if part.strip()]
        for part in parts:
            add(part)
        if len(parts) >= 2 and len(parts[-1].split()) <= 4:
            add(parts[-1])

    if "," in base and len(base.split(",")) <= 3:
        for part in (part.strip() for part in base.split(",")):
            add(part)

    return aliases[:5]


def _strict_rumble_candidate(raw_url: str) -> tuple[str, str] | None:
    """Return a strict canonical Rumble channel URL or ``None``.

    Only channel pages are accepted. This intentionally rejects generic
    root slugs and all non-channel Rumble paths.
    """

    candidate = canonicalize_channel_url(raw_url)
    if candidate is None or candidate.platform != "rumble":
        return None

    parsed = urlsplit(candidate.channel_url)
    path_parts = [part for part in parsed.path.split("/") if part]

    if len(path_parts) != 2 or path_parts[0] not in {"c", "user"}:
        return None
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{1,127}", path_parts[1] or ""):
        return None
    return candidate.channel_url, path_parts[1]


def _extract_rumble_url_from_cell(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None

    for candidate in extract_supported_channel_urls(text):
        strict = _strict_rumble_candidate(candidate.channel_url)
        if strict is not None:
            return strict[0]

    strict = _strict_rumble_candidate(text)
    if strict is not None:
        return strict[0]

    return None


def _find_header_maps(ws: Worksheet) -> SheetHeaderMap | None:
    """Locate a row containing both required headers."""

    row_count = ws.max_row or 0
    col_count = ws.max_column or 0
    if row_count < 1 or col_count < 1:
        return None

    for row_idx in range(1, min(row_count, 25) + 1):
        influencer_col = None
        rumble_col = None
        for col_idx in range(1, min(col_count, 120) + 1):
            value = ws.cell(row_idx, col_idx).value
            header = _normalize(str(value)) if value is not None else ""
            if header == "Influencer Name":
                influencer_col = col_idx
            elif header == "Rumble URL":
                rumble_col = col_idx
        if influencer_col and rumble_col:
            return SheetHeaderMap(
                sheet_name=ws.title,
                header_row=row_idx,
                influencer_col=influencer_col,
                rumble_col=rumble_col,
            )
    return None


def _iter_source_rows(ws: Worksheet, header_map: SheetHeaderMap) -> Iterable[SourceRow]:
    row_count = ws.max_row or header_map.header_row
    for row_idx in range(header_map.header_row + 1, row_count + 1):
        influencer = _normalize(str(ws.cell(row_idx, header_map.influencer_col).value or ""))
        if not influencer:
            continue
        rumble_raw = ws.cell(row_idx, header_map.rumble_col).value
        yield SourceRow(
            sheet_name=ws.title,
            excel_row=row_idx,
            influencer_name=influencer,
            rumble_url_raw=_normalize(str(rumble_raw or "")),
        )


def _search_serper(query: str, num: int = DEFAULT_SEARCH_LIMIT) -> list[dict[str, object]]:
    last_exc: Exception | None = None
    for attempt in range(1, SERPER_MAX_ATTEMPTS + 1):
        try:
            with httpx.Client(timeout=25.0) as client:
                response = client.post(
                    SERPER_URL,
                    headers={
                        "X-API-KEY": scraper_settings.serp_api_key,
                        "Content-Type": "application/json",
                    },
                    json={"q": query, "num": num},
                )
                if response.status_code in SERPER_RETRY_CODES:
                    response.raise_for_status()
                response.raise_for_status()
                payload = response.json()
                organic = payload.get("organic", []) if isinstance(payload, dict) else []
                return organic if isinstance(organic, list) else []
        except (httpx.HTTPError, ValueError) as exc:
            last_exc = exc
            if attempt >= SERPER_MAX_ATTEMPTS:
                break
            time.sleep(SERPER_RETRY_DELAY_SECONDS * attempt)
    if last_exc is not None:
        raise last_exc
    return []


def _candidate_queries(name: str) -> list[str]:
    aliases = _split_aliases(name)
    queries: list[str] = []
    for alias in aliases:
        queries.extend(
            [
                f'site:rumble.com/c/ "{alias}"',
                f'site:rumble.com/user/ "{alias}"',
                f'site:rumble.com "{alias}" rumble',
            ]
        )
    seen: set[str] = set()
    deduped: list[str] = []
    for query in queries:
        if query not in seen:
            seen.add(query)
            deduped.append(query)
    return deduped


def _extract_candidates_from_result(item: dict[str, object]) -> list[str]:
    values: list[str] = []
    for key in ("link", "title", "snippet", "displayedLink"):
        value = item.get(key)
        if isinstance(value, str) and value.strip():
            values.append(value.strip())

    urls: dict[str, str] = {}
    for value in values:
        strict = _strict_rumble_candidate(value)
        if strict is not None:
            urls[strict[0]] = strict[0]

    blob = " ".join(values)
    for candidate in extract_supported_channel_urls(blob):
        strict = _strict_rumble_candidate(candidate.channel_url)
        if strict is not None:
            urls[strict[0]] = strict[0]
    return list(urls.values())


def _score_candidate(
    candidate_url: str,
    influencer_name: str,
    query: str,
    item: dict[str, object],
    verified_title: str | None = None,
) -> float:
    score = 0.48
    normalized_name = _normalize_for_match(influencer_name)
    normalized_title = _normalize_for_match(str(item.get("title") or ""))
    normalized_snippet = _normalize_for_match(str(item.get("snippet") or ""))
    normalized_candidate = _normalize_for_match(candidate_url)

    parsed = urlsplit(candidate_url)
    path_slug = "".join(part for part in parsed.path.split("/") if part)
    normalized_path = _normalize_for_match(path_slug)

    if normalized_name and normalized_name == normalized_path:
        score += 0.30
    elif normalized_name and normalized_name in normalized_path:
        score += 0.20

    if normalized_name and normalized_name in normalized_title:
        score += 0.18
    elif normalized_name and normalized_name in normalized_snippet:
        score += 0.10

    if normalized_candidate and normalized_candidate in _normalize_for_match(str(item.get("link") or "")):
        score += 0.08

    if "site:rumble.com/c/" in query or "site:rumble.com/user/" in query:
        score += 0.05

    if verified_title:
        verified_norm = _normalize_for_match(verified_title)
        if normalized_name and normalized_name == verified_norm:
            score += 0.12
        elif normalized_name and normalized_name in verified_norm:
            score += 0.06

    return round(min(score, 0.99), 2)


def _extract_title_and_canonical(html: str) -> tuple[str | None, str | None]:
    title_match = re.search(r"<title[^>]*>(.*?)</title>", html, re.IGNORECASE | re.DOTALL)
    title = None
    if title_match:
        title = _normalize(re.sub(r"\s+", " ", title_match.group(1)))

    canonical_match = re.search(
        r'<link[^>]+rel=["\']canonical["\'][^>]+href=["\']([^"\']+)["\']',
        html,
        re.IGNORECASE,
    )
    if canonical_match:
        return title, canonical_match.group(1).strip()

    og_url_match = re.search(
        r'<meta[^>]+property=["\']og:url["\'][^>]+content=["\']([^"\']+)["\']',
        html,
        re.IGNORECASE,
    )
    if og_url_match:
        return title, og_url_match.group(1).strip()

    return title, None


def _verify_candidate(candidate_url: str) -> tuple[str | None, str | None]:
    try:
        with httpx.Client(timeout=20.0, follow_redirects=True) as client:
            response = client.get(candidate_url, headers={"User-Agent": "Mozilla/5.0"})
            response.raise_for_status()
            return _extract_title_and_canonical(response.text)
    except httpx.HTTPError:
        return None, None


@lru_cache(maxsize=2048)
def _resolve_rumble_from_serper_cached(name: str) -> tuple[RumbleCandidate, ...]:
    candidates: dict[str, RumbleCandidate] = {}
    seen_queries: set[str] = set()

    for query in _candidate_queries(name):
        if query in seen_queries:
            continue
        seen_queries.add(query)
        try:
            results = _search_serper(query=query, num=DEFAULT_SEARCH_LIMIT)
        except Exception:
            continue

        for item in results:
            if not isinstance(item, dict):
                continue
            for candidate_url in _extract_candidates_from_result(item):
                if candidate_url in candidates:
                    continue
                candidates[candidate_url] = RumbleCandidate(
                    channel_url=candidate_url,
                    confidence=_score_candidate(candidate_url, name, query, item),
                    source="serper",
                    query=query,
                    title=str(item.get("title") or "") or None,
                    snippet=str(item.get("snippet") or "") or None,
                    display_link=str(item.get("displayedLink") or "") or None,
                )

    ordered = sorted(candidates.values(), key=lambda candidate: candidate.confidence, reverse=True)
    return tuple(ordered)


def _best_candidate(name: str) -> RumbleCandidate | None:
    candidates = list(_resolve_rumble_from_serper_cached(name))
    if not candidates:
        return None

    top = candidates[0]
    if len(candidates) > 1 and (top.confidence - candidates[1].confidence) < 0.08:
        return None
    if top.confidence < DEFAULT_MIN_CONFIDENCE:
        return None

    verified_title, verified_url = _verify_candidate(top.channel_url)
    strict_verified = _strict_rumble_candidate(verified_url or "") if verified_url else None
    if strict_verified is not None:
        verified_url = strict_verified[0]
    if verified_title or verified_url:
        top = RumbleCandidate(
            channel_url=verified_url or top.channel_url,
            confidence=_score_candidate(
                verified_url or top.channel_url,
                name,
                top.query,
                {"title": top.title or "", "snippet": top.snippet or "", "link": top.channel_url},
                verified_title=verified_title,
            ),
            source="serper+verified",
            query=top.query,
            title=top.title,
            snippet=top.snippet,
            display_link=top.display_link,
            verified_title=verified_title,
            verified_canonical_url=verified_url,
        )

    if top.confidence < DEFAULT_MIN_CONFIDENCE:
        return None
    return top


def _dedupe_existing_channel(client, channel_url: str) -> dict[str, object] | None:
    try:
        result = (
            client.table("channels")
            .select("id,name,channel_url,secondary_urls")
            .eq("channel_url", channel_url)
            .limit(1)
            .execute()
        )
        rows = getattr(result, "data", None) or []
        if rows:
            return rows[0]
    except APIError:
        pass

    try:
        result = (
            client.table("channels")
            .select("id,name,channel_url,secondary_urls")
            .contains("secondary_urls", [channel_url])
            .limit(1)
            .execute()
        )
        rows = getattr(result, "data", None) or []
        if rows:
            return rows[0]
    except Exception:
        pass

    return None


def _insert_channel(client, *, name: str, channel_url: str, source: str, confidence: float) -> str:
    if _dedupe_existing_channel(client, channel_url) is not None:
        return "duplicate"

    now = _utc_now()
    payload = {
        "platform": "rumble",
        "channel_url": channel_url,
        "name": name,
        "description": "",
        "contact_info": [],
        "niche_tags": [],
        "is_active": True,
        "has_been_scraped": False,
        "discovery_source": source,
        "last_discovery_source": source,
        "discovery_status": "new",
        "discovery_confidence": confidence,
        "discovery_evidence_count": 1,
        "discovery_last_seen_at": now,
        "discovered_at": now,
        "updated_at": now,
    }
    client.table("channels").insert(payload).execute()
    return "inserted"


def _load_source_rows(workbook_path: Path, sheet_names: list[str] | None = None) -> tuple[list[SourceRow], list[SheetHeaderMap]]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    header_maps: list[SheetHeaderMap] = []
    rows: list[SourceRow] = []

    if sheet_names:
        normalized_targets = {_normalize(name): name for name in sheet_names}
        target_sheet_names = [
            sheet_name
            for sheet_name in workbook.sheetnames
            if _normalize(sheet_name) in normalized_targets
        ]
    else:
        target_sheet_names = list(workbook.sheetnames)
    for sheet_name in target_sheet_names:
        ws = workbook[sheet_name]
        header_map = _find_header_maps(ws)
        if header_map is None:
            continue
        header_maps.append(header_map)
        rows.extend(list(_iter_source_rows(ws, header_map)))

    return rows, header_maps


def _choose_rumble_url(row: SourceRow, min_confidence: float) -> tuple[str | None, str, float | None, str | None]:
    """Return ``(channel_url, resolution_source, confidence, error)``."""

    direct_url = _extract_rumble_url_from_cell(row.rumble_url_raw)
    if direct_url:
        return direct_url, "excel", 1.0, None

    best = _best_candidate(row.influencer_name)
    if best is None:
        return None, "serper", None, "no_high_confidence_match"

    if best.confidence < min_confidence:
        return None, best.source, best.confidence, "below_threshold"

    return best.verified_canonical_url or best.channel_url, best.source, best.confidence, None


def _log_row_resolution(
    index: int,
    total: int,
    row: SourceRow,
    chosen_url: str | None,
    source: str,
    confidence: float | None,
    error: str | None,
    action: str | None = None,
) -> None:
    prefix = f"[{index}/{total}] {row.sheet_name}!{row.excel_row} {row.influencer_name}"
    if chosen_url is None:
        if error == "below_threshold":
            print(f"{prefix} -> SERP skipped (confidence {confidence:.2f} below threshold)", flush=True)
        elif error == "no_high_confidence_match":
            print(f"{prefix} -> unresolved (no high-confidence SERP match)", flush=True)
        else:
            print(f"{prefix} -> unresolved", flush=True)
        return

    if source == "excel":
        print(f"{prefix} -> DIRECT URL {chosen_url}", flush=True)
    else:
        confidence_text = f"{confidence:.2f}" if confidence is not None else "n/a"
        print(f"{prefix} -> SERP {confidence_text} {chosen_url}", flush=True)

    if action:
        print(f"{prefix} -> {action}", flush=True)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Import Influencer Name / Rumble URL rows from an Excel workbook into Supabase."
    )
    parser.add_argument(
        "--xlsx",
        default="scraper/GoldCo Master Influencer List 2025-26 V1.xlsx",
        help="Path to the source Excel workbook",
    )
    parser.add_argument(
        "--sheet",
        action="append",
        default=None,
        help="Optional sheet name to restrict processing; repeat to include multiple sheets",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=0,
        help="Optional row limit for testing",
    )
    parser.add_argument(
        "--min-confidence",
        type=float,
        default=DEFAULT_MIN_CONFIDENCE,
        help="Minimum Serper confidence required to insert a resolved row",
    )
    parser.add_argument(
        "--write",
        action="store_true",
        help="Insert net-new channels into Supabase (default is dry-run)",
    )
    parser.add_argument(
        "--output",
        default="scraper/output/goldco_xlsx_rumble_import_report.json",
        help="JSON report path",
    )
    args = parser.parse_args()

    workbook_path = Path(args.xlsx)
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    source_rows, header_maps = _load_source_rows(workbook_path, args.sheet)
    if args.limit > 0:
        source_rows = source_rows[: args.limit]

    client = get_supabase_client() if args.write else None

    summary = {
        "mode": "write" if args.write else "dry_run",
        "workbook": str(workbook_path),
        "sheets_scanned": sorted({header_map.sheet_name for header_map in header_maps}),
        "rows_found": len(source_rows),
        "rows_with_excel_url": 0,
        "rows_needing_serper": 0,
        "resolved": 0,
        "inserted": 0,
        "duplicates": 0,
        "unresolved": 0,
        "below_threshold": 0,
        "candidate_urls": 0,
        "generated_at": _utc_now(),
    }
    report_rows: list[dict[str, object]] = []

    for index, row in enumerate(source_rows, start=1):
        direct_url = _extract_rumble_url_from_cell(row.rumble_url_raw)
        if direct_url:
            summary["rows_with_excel_url"] += 1
        else:
            summary["rows_needing_serper"] += 1

        chosen_url, source, confidence, error = _choose_rumble_url(row, args.min_confidence)
        candidate_details: dict[str, object] | None = None
        action = None

        if chosen_url is None:
            summary["unresolved"] += 1
            if error == "below_threshold":
                summary["below_threshold"] += 1
        else:
            summary["resolved"] += 1
            if source != "excel":
                best = _best_candidate(row.influencer_name)
                if best is not None:
                    candidate_details = {
                        "channel_url": best.channel_url,
                        "confidence": best.confidence,
                        "source": best.source,
                        "query": best.query,
                        "title": best.title,
                        "snippet": best.snippet,
                        "display_link": best.display_link,
                        "verified_title": best.verified_title,
                        "verified_canonical_url": best.verified_canonical_url,
                    }
                    summary["candidate_urls"] += 1

            if args.write and client is not None:
                status = _insert_channel(
                    client,
                    name=row.influencer_name,
                    channel_url=chosen_url,
                    source="xlsx_rumble_import" if source == "excel" else "xlsx_rumble_serper_import",
                    confidence=float(confidence or 0.0),
                )
                if status == "inserted":
                    summary["inserted"] += 1
                elif status == "duplicate":
                    summary["duplicates"] += 1
                action = status
            else:
                action = "candidate" if source == "excel" else "serper_candidate"

        _log_row_resolution(
            index=index,
            total=len(source_rows),
            row=row,
            chosen_url=chosen_url,
            source=source,
            confidence=confidence,
            error=error,
            action=action,
        )

        report_rows.append(
            {
                "sheet_name": row.sheet_name,
                "excel_row": row.excel_row,
                "influencer_name": row.influencer_name,
                "rumble_url_raw": row.rumble_url_raw,
                "resolved_url": chosen_url,
                "resolution_source": source,
                "confidence": confidence,
                "error": error,
                "candidate": candidate_details,
                "action": action,
            }
        )

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps({"summary": summary, "rows": report_rows}, indent=2), encoding="utf-8")
    print(json.dumps(summary, indent=2))
    print(f"Report written: {output_path}")


if __name__ == "__main__":
    main()
